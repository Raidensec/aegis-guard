from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


SRC = Path(r"C:\Users\ENG_AHMAD_ESSAM\Desktop\AI_SOC_Book_v3_polished.docx")
OUT = Path(r"C:\Users\ENG_AHMAD_ESSAM\Desktop\AI_SOC_Book_v4_baseera_format.docx")


def set_font(run, name="Times New Roman", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for key in ["ascii", "hAnsi", "eastAsia", "cs"]:
        r_fonts.set(qn(f"w:{key}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_para_spacing(p, before=0, after=6, line=1.15):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def clear_paragraph(p):
    p.clear()
    return p


def delete_paragraph(p):
    el = p._element
    el.getparent().remove(el)
    p._p = p._element = None


def paragraph_before(p, text="", style=None):
    new_p = OxmlElement("w:p")
    p._p.addprevious(new_p)
    para = Paragraph(new_p, p._parent)
    if style:
        para.style = style
    if text:
        para.add_run(text)
    set_para_spacing(para)
    return para


def paragraph_after(p, text="", style=None):
    new_p = OxmlElement("w:p")
    p._p.addnext(new_p)
    para = Paragraph(new_p, p._parent)
    if style:
        para.style = style
    if text:
        para.add_run(text)
    set_para_spacing(para)
    return para


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_font(run, "Calibri", 11, bold=True)


def normalize_caption_text(text):
    return re.sub(r"^(Figure|Table)\s+(\d+):\s*", r"\1 \2 ", text.strip())


def add_front_entry(anchor, title, page="0", indent=0, size=14, bold=False):
    p = paragraph_before(anchor)
    p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.tab_stops.add_tab_stop(
        Inches(7.05 - indent),
        WD_TAB_ALIGNMENT.RIGHT,
        WD_TAB_LEADER.DOTS,
    )
    r = p.add_run(title)
    set_font(r, size=size, bold=bold)
    r2 = p.add_run("\t" + str(page))
    set_font(r2, size=size, bold=bold)
    set_para_spacing(p, after=8 if size >= 14 else 4, line=1.0)
    return p


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill="FFFFFF"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="808080", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), size)
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def main():
    shutil.copy2(SRC, OUT)
    doc = Document(OUT)

    if "Caption" not in [s.name for s in doc.styles]:
        doc.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)

    for section in doc.sections:
        section.start_type = WD_SECTION_START.NEW_PAGE
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.35)
        for header in (section.header, section.first_page_header, section.even_page_header):
            header.is_linked_to_previous = False
            for p in header.paragraphs:
                p.clear()
        footer = section.footer
        for p in footer.paragraphs:
            p.clear()
        fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_page_field(fp)
        pipe = fp.add_run(" | ")
        set_font(pipe, "Calibri", 11, bold=True)
        page = fp.add_run("P a g e")
        set_font(page, "Calibri", 11)

    # Baseera-like style system: Times New Roman, black hierarchy.
    style_specs = {
        "Normal": (12, False, False),
        "Heading 1": (16, True, False),
        "Heading 2": (16, True, False),
        "Heading 3": (14, True, False),
        "Caption": (9, False, True),
    }
    for name, (size, bold, italic) in style_specs.items():
        s = doc.styles[name]
        s.font.name = "Times New Roman"
        s._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        s._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        s.font.size = Pt(size)
        s.font.bold = bold
        s.font.italic = italic
        s.font.color.rgb = RGBColor(0, 0, 0)

    for p in doc.paragraphs:
        text = p.text.strip()
        if p.style.name == "Normal":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if len(text) > 80 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Inches(0.25) if len(text) > 80 else None
            set_para_spacing(p, before=0, after=6, line=1.15)
        elif p.style.name == "Heading 1":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_para_spacing(p, before=12, after=8, line=1.0)
        elif p.style.name == "Heading 2":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_para_spacing(p, before=12, after=6, line=1.0)
        elif p.style.name == "Heading 3":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_para_spacing(p, before=10, after=4, line=1.0)
        elif p.style.name == "Caption":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_para_spacing(p, before=4, after=8, line=1.0)
        for run in p.runs:
            if p.style.name == "Caption":
                set_font(run, size=9, italic=True)
            elif p.style.name.startswith("Heading"):
                set_font(run, size=p.style.font.size.pt if p.style.font.size else 16, bold=True)
            else:
                set_font(run, size=12)

    # Remove generated front matter from previous version.
    toc_idx = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "Table of Contents")
    chapter_hits = [i for i, p in enumerate(doc.paragraphs) if p.text.strip().startswith("CHAPTER (1)")]
    body_chapter_idx = chapter_hits[-1]
    for p in list(doc.paragraphs[toc_idx:body_chapter_idx]):
        delete_paragraph(p)

    # Normalize captions to Baseera's "Figure N Title" / "Table N Title" form.
    for p in doc.paragraphs:
        text = p.text.strip()
        if re.match(r"^(Figure|Table)\s+\d+:", text):
            p.clear()
            run = p.add_run(normalize_caption_text(text))
            if text.startswith("Figure"):
                p.style = "Caption"
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_font(run, size=9, italic=True)
            else:
                p.style = "Normal"
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                set_font(run, size=12, bold=False)

    # Chapter divider pages like the Baseera reference.
    for p in list(doc.paragraphs):
        text = p.text.strip()
        match = re.match(r"^CHAPTER\s+\((\d+)\)\s+(.+)$", text)
        if not match:
            continue
        number, title = match.groups()
        title = title.strip()
        p.clear()
        r1 = p.add_run(f"CHAPTER ({number})")
        r1.add_break()
        r2 = p.add_run(title)
        set_font(r1, size=28, bold=True)
        set_font(r2, size=28, bold=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.style = "Heading 1"
        set_para_spacing(p, before=185, after=0, line=1.2)
        before = paragraph_before(p)
        before.add_run().add_break(WD_BREAK.PAGE)
        after = paragraph_after(p)
        after.add_run().add_break(WD_BREAK.PAGE)

    # Tables: simple Word report tables, no blue fills.
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for ri, row in enumerate(table.rows):
            if ri == 0:
                tr_pr = row._tr.get_or_add_trPr()
                if tr_pr.find(qn("w:tblHeader")) is None:
                    tbl_header = OxmlElement("w:tblHeader")
                    tbl_header.set(qn("w:val"), "true")
                    tr_pr.append(tbl_header)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                set_cell_borders(cell)
                set_cell_shading(cell, "FFFFFF")
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.first_line_indent = None
                    set_para_spacing(p, before=0, after=0, line=1.0)
                    for run in p.runs:
                        set_font(run, size=10 if len(table.columns) >= 5 else 11, bold=(ri == 0))

    # Gather entries from body and insert placeholder front matter. A second pass
    # rebuilds these with true Word page numbers.
    body_chapter = next(p for p in doc.paragraphs if p.text.strip().startswith("CHAPTER (1)"))
    headings = []
    figures = []
    tables = []
    body_started = False
    for p in doc.paragraphs:
        text = p.text.strip()
        if text.startswith("CHAPTER (1)"):
            body_started = True
        if not body_started or not text:
            continue
        if p.style.name in ["Heading 1", "Heading 2", "Heading 3"] or re.match(r"^\d+\.\d+", text):
            headings.append((text.replace("\n", " "), p.style.name))
        elif text.startswith("Figure "):
            figures.append(text)
        elif text.startswith("Table "):
            tables.append(text)

    front = [("heading", "Table of Contents", None, 0, 16, False)]
    for text, style in headings:
        if text.startswith("CHAPTER"):
            front.append(("entry", text, "0", 0, 14, True))
        elif style == "Heading 3":
            front.append(("entry", text, "0", 0.45, 12, False))
        elif style == "Heading 2" or re.match(r"^\d+\.\d+", text):
            front.append(("entry", text, "0", 0, 14, False))
    front += [("break", None, None, 0, 0, False), ("heading", "Table of Figures:", None, 0, 16, False)]
    front += [("entry", item, "0", 0, 12, False) for item in figures]
    front += [("break", None, None, 0, 0, False), ("heading", "Table of Tables", None, 0, 16, False)]
    front += [("entry", item, "0", 0, 12, False) for item in tables]
    front += [("break", None, None, 0, 0, False)]

    for kind, text, page, indent, size, bold in front:
        if kind == "break":
            bp = paragraph_before(body_chapter)
            bp.add_run().add_break(WD_BREAK.PAGE)
        elif kind == "heading":
            hp = paragraph_before(body_chapter, text, "Heading 1")
            hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_para_spacing(hp, before=12, after=18, line=1.0)
            for run in hp.runs:
                set_font(run, "Cambria" if text == "Table of Contents" else "Times New Roman", size=16, bold=False)
        else:
            add_front_entry(body_chapter, text, page, indent=indent, size=size, bold=bold)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
